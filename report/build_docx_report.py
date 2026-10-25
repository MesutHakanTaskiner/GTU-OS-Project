from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report" / "report_source.md"
OUTPUT = ROOT / "deliverables" / "CSE513_Assignment1_Report.docx"

BLACK = "000000"
NAVY = "17365D"
PALE_BLUE = "EAF1F8"
LIGHT_GREY = "D9D9D9"
CODE_GREY = "F3F3F3"


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=85, bottom=85, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), LIGHT_GREY)


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, url: str, bold: bool = False) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    if bold:
        run_properties.append(OxmlElement("w:b"))
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_TOKEN = re.compile(
    r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\)|https?://[^\s)]+)"
)


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            inner = token[2:-2]
            bold_cursor = 0
            for linked in re.finditer(r"\[([^\]]+)\]\((https?://[^)]+)\)", inner):
                if linked.start() > bold_cursor:
                    paragraph.add_run(inner[bold_cursor : linked.start()]).bold = True
                add_hyperlink(paragraph, linked.group(1), linked.group(2), bold=True)
                bold_cursor = linked.end()
            if bold_cursor < len(inner):
                paragraph.add_run(inner[bold_cursor:]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.size = Pt(9.5)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            add_hyperlink(paragraph, token, token)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def clean_heading(text: str) -> str:
    text = text.replace("—", " ").replace("–", " ").replace("-", " ")
    text = text.replace("=", " ").replace(",", " ").replace(":", " ")
    text = re.sub(r"(?<=\d)\.(?=\d)", " ", text)
    text = re.sub(r"(?<=\d)\.(?=\s)", "", text)
    return re.sub(r"\s+", " ", text).strip()


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [[item.strip() for item in line.strip().strip("|").split("|")] for line in lines]
    rows = [rows[0], *rows[2:]]
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    column_count = len(rows[0])
    if column_count == 8:
        widths = [0.82, 0.68, 0.62, 0.88, 0.48, 0.48, 0.48, 0.48]
        font_size = 8.0
    elif column_count == 4:
        widths = [1.9, 1.45, 1.45, 1.45]
        font_size = 9.0
    elif column_count == 3:
        widths = [2.4, 1.45, 1.45]
        font_size = 9.2
    else:
        widths = [6.7 / column_count] * column_count
        font_size = 8.8

    for row_index, values in enumerate(rows):
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_properties.append(cant_split)
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            if row_index < len(rows) - 1:
                paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if row_index == 0:
                shade_cell(cell, NAVY)
            elif row_index % 2 == 0:
                shade_cell(cell, PALE_BLUE)

    following = document.add_paragraph()
    following.paragraph_format.space_after = Pt(1)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 90, 90)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = document.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(9)
    title.paragraph_format.keep_with_next = True
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 14.5, 11, 5),
        ("Heading 2", 12.0, 8, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    list_bullet = document.styles["List Bullet"]
    list_bullet.font.name = "Calibri"
    list_bullet.font.size = Pt(10.5)
    list_bullet.font.color.rgb = RGBColor(0, 0, 0)
    list_bullet.paragraph_format.space_after = Pt(2)


def build_document(source_text: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.75)
    configure_styles(document)
    add_page_number(section.footer.paragraphs[0])

    lines = source_text.splitlines()
    index = 0
    in_code_block = False
    while index < len(lines):
        stripped = lines[index].strip()
        if stripped == "<!-- PAGE_BREAK -->":
            document.add_page_break()
            index += 1
            continue
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            index += 1
            continue
        if in_code_block:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.35)
            paragraph.paragraph_format.right_indent = Cm(0.35)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(lines[index])
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.size = Pt(9)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), CODE_GREY)
            paragraph._p.get_or_add_pPr().append(shade)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, clean_heading(stripped[2:]))
            index += 1
            continue
        if stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, clean_heading(stripped[3:]))
            index += 1
            continue
        if stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, clean_heading(stripped[4:]))
            index += 1
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue

        paragraph_parts = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "|", "- ", "```", "<!-- PAGE_BREAK"))
            ):
                break
            paragraph_parts.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_parts))

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.widow_control = True
        for run in paragraph.runs:
            if run.font.color.rgb is None:
                run.font.color.rgb = RGBColor(0, 0, 0)

    core = document.core_properties
    core.title = "CSE513 Assignment 1 CPU Scheduling Simulator"
    core.subject = "Academic report for a browser-based CPU scheduling simulator"
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.keywords = "CPU scheduling, FCFS, SJF, SRTF, Round Robin, Streamlit"
    return document


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document(SOURCE.read_text(encoding="utf-8"))
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
